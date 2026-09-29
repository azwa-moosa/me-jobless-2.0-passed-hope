"""Build rtm.csv, rtm.xlsx and RTM markdown from rtm_data.py.  Usage: python build_rtm.py"""
import csv
from collections import OrderedDict
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import rtm_data as d

HERE = __file__.rsplit("/", 1)[0] if "/" in __file__ else "."

# CSV (diff-friendly, version-controlled)
with open(f"{HERE}/rtm.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(d.COLUMNS)
    w.writerows(d.R)

# XLSX
wb = Workbook()
ws = wb.active
ws.title = "RTM"
F = "Arial"
thin = Side(style="thin", color="BFBFBF")
hdr_fill = PatternFill("solid", fgColor="1F3864")
ws.append(d.COLUMNS)
for row in d.R:
    ws.append(list(row))
widths = [9, 18, 42, 32, 22, 20, 24, 30, 26, 28, 34, 24, 30, 8, 12, 10, 42, 14, 16]
for i, wdt in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = wdt
for c in ws[1]:
    c.font = Font(name=F, bold=True, color="FFFFFF", size=10)
    c.fill = hdr_fill
    c.alignment = Alignment(wrap_text=True, vertical="center")
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=F, size=9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
ws.freeze_panes = "D2"
ref = f"A1:{get_column_letter(len(d.COLUMNS))}{len(d.R) + 1}"
t = Table(displayName="RTM", ref=ref)
t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
ws.add_table(t)

# Summary sheet – live formulas over the RTM sheet
s = wb.create_sheet("Summary")
modules = list(OrderedDict.fromkeys(r[1] for r in d.R))
s["A1"] = "RTM coverage by module and priority (formulas over RTM sheet)"
s["A1"].font = Font(name=F, bold=True, size=12)
heads = ["Module", "P0", "P1", "P2", "Total"]
for j, h in enumerate(heads, 1):
    c = s.cell(row=3, column=j, value=h)
    c.font = Font(name=F, bold=True, color="FFFFFF")
    c.fill = hdr_fill
n = len(d.R) + 1
for i, m in enumerate(modules, 4):
    s.cell(row=i, column=1, value=m)
    for j, p in enumerate(["P0", "P1", "P2"], 2):
        s.cell(row=i, column=j, value=f'=COUNTIFS(RTM!$B$2:$B${n},$A{i},RTM!$N$2:$N${n},"{p}")')
    s.cell(row=i, column=5, value=f"=SUM(B{i}:D{i})")
last = 3 + len(modules)
tr = last + 1
s.cell(row=tr, column=1, value="Total")
for j in range(2, 6):
    col = get_column_letter(j)
    s.cell(row=tr, column=j, value=f"=SUM({col}4:{col}{last})")
for row in s.iter_rows(min_row=4, max_row=tr):
    for c in row:
        c.font = Font(name=F, bold=(c.row == tr))
s.column_dimensions["A"].width = 30
s.cell(row=tr + 2, column=1, value="Check: rows in RTM").font = Font(name=F)
s.cell(row=tr + 2, column=2, value=f"=COUNTA(RTM!$A$2:$A${n})").font = Font(name=F)

leg = wb.create_sheet("Legend")
rows = [
    ("Priority source", "BP §22 = stated in blueprint backlog; Derived = required to deliver a BP §22 story or blueprint control; Assumed = build-team proposal awaiting confirmation"),
    ("Phase", "Blueprint §21 phase numbers. '4b (proposed)' = resequencing proposal DR-43, not yet approved"),
    ("Open decisions", "DR-xx items in docs/00-planning/L-decisions-register.md"),
    ("Source of truth", "Edit docs/requirements/rtm_data.py then run build_rtm.py; do not edit this workbook directly"),
]
for i, (k, v) in enumerate(rows, 1):
    leg.cell(row=i, column=1, value=k).font = Font(name=F, bold=True)
    c = leg.cell(row=i, column=2, value=v)
    c.font = Font(name=F)
    c.alignment = Alignment(wrap_text=True)
leg.column_dimensions["A"].width = 18
leg.column_dimensions["B"].width = 110
wb.save(f"{HERE}/rtm.xlsx")

# Markdown (grouped by module; compact column subset + link to full CSV)
md = ["# B. Requirements Traceability Matrix", "",
      "Full 19-column matrix: `docs/requirements/rtm.csv` (source: `rtm_data.py`; Excel view: `rtm.xlsx`).",
      "This page shows the compact view. Priority source and decision references are explained in the Legend sheet.", "",
      f"**{len(d.R)} requirements** traced to the blueprint.", ""]
for m in modules:
    md += [f"## {m}", "", "| ID | Requirement | Roles | Service | Entities | Security | Audit | AI | Pri | Phase | Acceptance | BP | Decisions |",
           "|---|---|---|---|---|---|---|---|---|---|---|---|---|"]
    for r in d.R:
        if r[1] != m:
            continue
        cells = [r[0], r[2], r[4], r[6], r[7], r[10], r[11], r[12], f"{r[13]} ({r[14]})", r[15], r[16], r[17], r[18]]
        md.append("| " + " | ".join(str(x).replace("|", "/") for x in cells) + " |")
    md.append("")
with open(f"{HERE}/../00-planning/B-requirements-traceability-matrix.md", "w") as f:
    f.write("\n".join(md))
print("built", len(d.R))
